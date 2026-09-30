from io import BytesIO

from docx import Document
from fpdf import FPDF


# =========================================================
# TXT EXPORT
# =========================================================

def export_to_txt(text: str) -> bytes:
    """
    Convert document text to TXT bytes.
    """
    return text.encode("utf-8")


# =========================================================
# DOCX EXPORT
# =========================================================

def export_to_docx(text: str) -> bytes:
    """
    Convert document text to a Microsoft Word DOCX file.
    """

    document = Document()

    # Split document into lines
    lines = text.splitlines()

    for line in lines:

        line = line.rstrip()

        # Empty line
        if not line:
            document.add_paragraph()
            continue

        # Detect simple headings
        if (
            len(line) <= 100
            and line.upper() == line
        ):
            paragraph = document.add_paragraph()

            run = paragraph.add_run(line)

            run.bold = True

            continue

        # Normal paragraph
        document.add_paragraph(line)

    # Save DOCX into memory
    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


# =========================================================
# PDF EXPORT
# =========================================================

def export_to_pdf(text: str) -> bytes:
    """
    Convert document text to PDF bytes.
    """

    pdf = FPDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=15
    )

    pdf.add_page()

    pdf.set_font(
        "Arial",
        size=11
    )

    for line in text.splitlines():

        line = line.strip()

        # Empty line
        if not line:
            pdf.ln(5)
            continue

        # Standard Arial PDF font does not support
        # every Unicode character.
        safe_line = (
            line
            .replace("₹", "Rs.")
            .replace("–", "-")
            .replace("—", "-")
            .replace("“", '"')
            .replace("”", '"')
            .replace("’", "'")
            .replace("•", "-")
        )

        pdf.multi_cell(
            0,
            7,
            safe_line
        )

    # fpdf2 returns a bytearray
    return bytes(pdf.output())